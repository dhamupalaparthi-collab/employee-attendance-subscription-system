# filepath: read_reference.py
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# Read the reference document
ref_doc = Document(r"c:\Users\dhamodharam\Downloads\OBJECT ORIENTED PROGRAMMING PROJECT (1).docx")

# Print all paragraphs to understand structure
print("=== REFERENCE DOCUMENT STRUCTURE ===\n")

for i, para in enumerate(ref_doc.paragraphs):
    text = para.text.strip()
    if text:
        # Check if it has a style
        style_name = para.style.name if para.style else "Normal"
        print(f"[{i}] Style: {style_name}")
        print(f"    Text: {text[:100]}...")
        print()

print(f"\nTotal paragraphs: {len(ref_doc.paragraphs)}")
print(f"Total tables: {len(ref_doc.tables)}")

# Print table structures
for i, table in enumerate(ref_doc.tables):
    print(f"\n=== TABLE {i+1} ===")
    print(f"Rows: {len(table.rows)}, Columns: {len(table.columns)}")
    if table.rows:
        # Print first row (header)
        print("Header row:")
        for cell in table.rows[0].cells:
            print(f"  {cell.text[:50]}")